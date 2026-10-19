import csv
import sys
from io import BytesIO, StringIO
from unittest.mock import patch
from zipfile import ZipFile

from django.contrib.admin.models import LogEntry
from django.contrib.auth.models import Group, Permission, User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client, RequestFactory, SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.views.debug import ExceptionReporter
from openpyxl import Workbook

from .admission_results import HEADERS, MAX_UPLOAD_BYTES, identity_key, read_results, selection_token
from .models import AdmissionRelease, AdmissionResult, SiteSettings, SubmissionThrottle

ROW = [2027, 'Mokoena', 'Thabo James', 'St Mary’s Primary', 'Y']


def upload(rows=None, extension='csv', header=HEADERS):
    rows = [ROW] if rows is None else rows
    if extension == 'xlsx':
        wb = Workbook()
        wb.active.append(list(header))
        for row in rows:
            wb.active.append(row)
        stream = BytesIO()
        wb.save(stream)
        wb.close()
        data = stream.getvalue()
    else:
        stream = StringIO()
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)
        data = stream.getvalue().encode('utf-8')
    return SimpleUploadedFile(f'admitted.{extension}', data)


class AdmissionImportParsingTests(SimpleTestCase):
    def test_csv_and_excel_and_boarding_blank_or_n(self):
        for extension in ('csv', 'xlsx'):
            with self.subTest(extension=extension):
                rows = read_results(upload([ROW, [2027, 'Example', 'Anna', 'Example Primary', None],
                    [2027, 'Example', 'James', 'Example Primary', 'n']], extension), 2027)
                self.assertEqual([row['boarding'] for row in rows], [True, False, False])
                self.assertEqual(set(rows[0]), {'identity_key', 'surname', 'first_names', 'primary_school', 'boarding'})
                self.assertEqual(rows[0]['surname'], ROW[1])
                self.assertEqual(rows[0]['primary_school'], ROW[3])

    def test_normalization_matches_whitespace_case_and_apostrophes(self):
        self.assertEqual(identity_key(*ROW[:4]), identity_key(2027, ' MOKOENA ', 'thabo   james', "St Mary's Primary"))
        self.assertNotEqual(identity_key(*ROW[:4]), identity_key(2028, *ROW[1:4]))

    def test_rejects_bad_rows_without_including_student_names_in_errors(self):
        bad_rows = [[], [ROW, ROW], [[2028, *ROW[1:]]], [[2027, '', *ROW[2:]]],
            [[*ROW[:4], 'Yes']], [[*ROW[:4], 0]], [[*ROW[:2], 'A' * 121, *ROW[3:]]], [[*ROW, 'Extra']],
            [[2027, '=SUM(1)', *ROW[2:]]], [[2027, 'Mokoena', 'Thabo', '', 'Y']]]
        for rows in bad_rows:
            with self.subTest(rows=rows):
                with self.assertRaises(ValidationError) as error:
                    read_results(upload(rows), 2027)
                self.assertNotIn('Mokoena', str(error.exception))

    def test_rejects_wrong_headers_types_and_file_size(self):
        for file in [upload(header=('Incorrect',)),
                     SimpleUploadedFile('bad.xls', b'not excel'),
                     SimpleUploadedFile('bad.xlsx', b'not excel'),
                     SimpleUploadedFile('bad.csv', b'\xff'),
                     SimpleUploadedFile('large.csv', b'A' * (MAX_UPLOAD_BYTES + 1))]:
            with self.subTest(name=file.name), self.assertRaises(ValidationError):
                read_results(file, 2027)

    def test_excel_formulas_errors_multiple_sheets_and_xml_entities_rejected(self):
        for value in ('=1+1', '#DIV/0!'):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                read_results(upload([[2027, value, *ROW[2:]]], 'xlsx'), 2027)
        stream = BytesIO()
        wb = Workbook()
        wb.create_sheet('Extra')
        wb.save(stream)
        with self.assertRaises(ValidationError):
            read_results(SimpleUploadedFile('multiple.xlsx', stream.getvalue()), 2027)
        stream = BytesIO()
        with ZipFile(stream, 'w') as archive:
            archive.writestr('malicious.xml', '<!DOCTYPE a [<!ENTITY x "bad">]><a>&x;</a>')
        with self.assertRaises(ValidationError):
            read_results(SimpleUploadedFile('unsafe.xlsx', stream.getvalue()), 2027)

    def test_archive_expansion_and_row_caps(self):
        stream = BytesIO()
        with ZipFile(stream, 'w') as archive:
            archive.writestr('oversized.xml', 'A' * 100)
        with patch('school.admission_results.MAX_UNPACKED_BYTES', 50), self.assertRaises(ValidationError):
            read_results(SimpleUploadedFile('large.xlsx', stream.getvalue()), 2027)
        with patch('school.admission_results.MAX_ROWS', 1), self.assertRaises(ValidationError):
            read_results(upload([ROW, [2027, 'Example', 'Anna', 'Example Primary', 'N']]), 2027)


@override_settings(SECURE_SSL_REDIRECT=False, DEBUG=False, SITE_URL='https://school.example.test',
                   ALLOWED_HOSTS=['testserver'], STORAGES={
    'default': {'BACKEND': 'django.core.files.storage.InMemoryStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
})
class AdmissionResultsWorkflowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        SiteSettings.objects.create(demo_mode=False, admissions_open=False)
        call_command('setup_staff_groups', verbosity=0)
        cls.office = User.objects.create_user('results-office', is_staff=True)
        cls.office.groups.add(Group.objects.get(name='Admissions officers'))
        cls.editor = User.objects.create_user('results-editor', is_staff=True)
        cls.editor.groups.add(Group.objects.get(name='Content editors'))
        cls.release = AdmissionRelease.objects.create(year=2027, published=True)
        cls.result = AdmissionResult.objects.create(release=cls.release, identity_key=identity_key(*ROW[:4]),
            surname=ROW[1], first_names=ROW[2], primary_school=ROW[3], boarding=True)

    def search(self, **changes):
        data = {'surname': ROW[1], 'first_names': ROW[2]}
        data.update(changes)
        return self.client.post(reverse('admission-results'), data)

    def view_result(self, result=None, **changes):
        result = result or AdmissionResult.objects.select_related('release').get(identity_key=identity_key(*ROW[:4]))
        data = {'action': 'view', 'selection': selection_token(result)}
        data.update(changes)
        return self.client.post(reverse('admission-results'), data)

    def staff_action(self, action, **changes):
        self.release.refresh_from_db()
        data = {'year': self.release.year, 'revision': str(self.release.revision), 'action': action}
        data.update(changes)
        return self.client.post(reverse('staff-admission-results'), data)

    def staff_upload(self, rows=None, year=2027):
        return self.client.post(reverse('staff-admission-results'), {
            'action': 'upload', 'year': year, 'replace': 'on', 'spreadsheet': upload(rows)})

    def test_selected_student_shows_admission_and_boarding_without_confirmation_or_next_steps(self):
        response = self.view_result()
        self.assertContains(response, 'Admitted.')
        self.assertContains(response, 'Boarding: Offered')
        self.assertContains(response, ROW[1])
        self.assertContains(response, ROW[2])
        self.assertContains(response, ROW[3])
        self.assertNotContains(response, 'name="primary_school"')
        self.assertNotContains(response, 'accepting the offer')
        self.assertNotContains(response, 'next steps')
        self.assertNotContains(response, self.result.identity_key)
        self.assertNotIn('Location', response)
        self.assertFalse(any('result' in key for key in self.client.session.keys()))
        self.result.boarding = False
        self.result.save()
        self.assertContains(self.view_result(), 'Boarding: Not offered')

    def test_wrong_names_and_unpublished_results_share_generic_reply(self):
        for changes in ({'first_names': 'Missing'}, {'surname': 'Other'}, {'surname': 'Missing'}):
            response = self.search(**changes)
            self.assertContains(response, 'No matching students found')
            self.assertNotContains(response, 'Boarding:')
        SubmissionThrottle.objects.all().delete()
        self.release.published = False
        self.release.save()
        response = self.search()
        self.assertContains(response, 'No matching students found')
        self.assertNotContains(response, 'Boarding: Offered')

    def test_get_parameters_head_and_empty_forms_never_reveal_results(self):
        for method in ('get', 'head'):
            response = getattr(self.client, method)(reverse('admission-results'), {'surname': ROW[1]})
            self.assertEqual(response.status_code, 200)
            self.assertNotIn(b'Boarding:', response.content)
            self.assertIn('no-store', response['Cache-Control'])
            self.assertIn('noindex', response['X-Robots-Tag'])
            self.assertEqual(response['Referrer-Policy'], 'same-origin')
        response = self.search(surname='')
        self.assertNotContains(response, 'Boarding:')
        self.assertEqual(self.client.put(reverse('admission-results')).status_code, 405)

    @override_settings(ADMISSION_LOOKUP_RATE_PER_MINUTE=3)
    def test_invalid_csrf_attempts_are_limited_before_lookup(self):
        client = Client(enforce_csrf_checks=True)
        for _ in range(3):
            response = client.post(reverse('admission-results'), {})
            self.assertEqual(response.status_code, 403)
            self.assertIn('no-store', response['Cache-Control'])
        response = client.post(reverse('admission-results'), {})
        self.assertEqual(response.status_code, 429)
        self.assertIn('Retry-After', response)
        self.assertIn('noindex', response['X-Robots-Tag'])

    def test_csrf_verified_lookup_accepts_same_origin_http_https_and_https_referer(self):
        client = Client(enforce_csrf_checks=True)
        path = reverse('admission-results')
        for secure, headers in [
            (False, {'HTTP_ORIGIN': 'http://testserver'}),
            (True, {'HTTP_ORIGIN': 'https://testserver'}),
            (True, {'HTTP_REFERER': 'https://testserver' + path}),
        ]:
            with self.subTest(secure=secure, headers=headers):
                page = client.get(path, secure=secure)
                self.assertEqual(page['Referrer-Policy'], 'same-origin')
                response = client.post(path, {
                    'csrfmiddlewaretoken': client.cookies['csrftoken'].value,
                    'year': ROW[0], 'surname': ROW[1], 'first_names': ROW[2],
                    'primary_school': ROW[3], 'own_details': 'on',
                }, secure=secure, **headers)
                self.assertContains(response, 'Admitted students ·')
                self.assertNotContains(response, 'Boarding: Offered')

    def test_csrf_still_rejects_null_and_foreign_origins_with_valid_token(self):
        client = Client(enforce_csrf_checks=True)
        path = reverse('admission-results')
        client.get(path)
        for origin in ('null', 'https://untrusted.example.test'):
            with self.subTest(origin=origin):
                response = client.post(path, {
                    'csrfmiddlewaretoken': client.cookies['csrftoken'].value,
                    'year': ROW[0], 'surname': ROW[1], 'first_names': ROW[2],
                    'primary_school': ROW[3], 'own_details': 'on',
                }, HTTP_ORIGIN=origin)
                self.assertEqual(response.status_code, 403)
                self.assertEqual(response['Referrer-Policy'], 'same-origin')
                self.assertIn('no-store', response['Cache-Control'])
                self.assertNotContains(response, 'Boarding:', status_code=403)

    def test_staff_publication_accepts_csrf_verified_same_origin_post(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.office)
        path = reverse('staff-admission-results')
        page = client.get(path)
        self.assertEqual(page['Referrer-Policy'], 'same-origin')
        response = client.post(path, {
            'csrfmiddlewaretoken': client.cookies['csrftoken'].value,
            'year': self.release.year, 'revision': str(self.release.revision),
            'action': 'unpublish',
        }, HTTP_ORIGIN='http://testserver')
        self.assertEqual(response.status_code, 302)
        self.release.refresh_from_db()
        self.assertFalse(self.release.published)

    @override_settings(ADMISSION_LOOKUP_RATE_PER_HOUR=2)
    def test_hourly_ip_limit_includes_invalid_searches(self):
        for _ in range(2):
            self.assertEqual(self.search(surname='').status_code, 200)
        self.assertEqual(self.view_result().status_code, 429)

    @override_settings(ADMISSION_LOOKUP_RATE_PER_DAY=1)
    def test_daily_ip_limit(self):
        self.assertEqual(self.search().status_code, 200)
        self.assertEqual(self.search().status_code, 429)

    @override_settings(ADMISSION_DETAIL_RATE_PER_DAY=20, ADMISSION_LOOKUP_RATE_PER_MINUTE=3)
    def test_identity_limit_across_addresses_and_forwarded_headers_not_trusted(self):
        key = self.result.identity_key
        from .submissions import consume_limit
        for _ in range(20):
            self.assertEqual(consume_limit(key, 'admission-identity', 20, 86400), 0)
        self.assertEqual(self.view_result().status_code, 429)
        self.assertFalse(SubmissionThrottle.objects.filter(key__contains='Mokoena').exists())
        SubmissionThrottle.objects.all().delete()
        for n in range(3):
            response = self.client.post(reverse('admission-results'), {}, HTTP_X_FORWARDED_FOR=f'192.0.2.{n}')
            self.assertEqual(response.status_code, 200)
        self.assertEqual(self.search().status_code, 429)

    def test_search_needs_no_year_and_displays_latest_published_intake_only(self):
        old = AdmissionRelease.objects.create(year=2026, published=True)
        old_result = AdmissionResult.objects.create(release=old, surname=ROW[1], first_names='Historical Learner',
            primary_school='Old Primary', identity_key=identity_key(2026, ROW[1], 'Historical Learner', 'Old Primary'))
        future = AdmissionRelease.objects.create(year=2028, published=False)
        AdmissionResult.objects.create(release=future, surname=ROW[1], first_names='Future Learner',
            primary_school='New Primary', identity_key=identity_key(2028, ROW[1], 'Future Learner', 'New Primary'))
        page = self.client.get(reverse('admission-results'))
        self.assertContains(page, 'Find your name · 2027')
        self.assertNotContains(page, 'name="year"')
        response = self.search(first_names='', year=2026)
        self.assertContains(response, 'Mokoena, Thabo James')
        self.assertContains(response, 'Admitted students · 2027')
        self.assertNotContains(response, 'Historical Learner')
        self.assertNotContains(response, 'Future Learner')
        self.assertNotContains(self.view_result(old_result), 'Boarding:')

    def test_old_selection_stops_working_when_new_intake_is_published(self):
        token = selection_token(self.result)
        current = AdmissionRelease.objects.create(year=2028, published=True)
        AdmissionResult.objects.create(release=current, surname=ROW[1], first_names='Current Learner',
            primary_school='Current Primary', identity_key=identity_key(2028, ROW[1], 'Current Learner', 'Current Primary'))
        self.assertNotContains(self.view_result(selection=token), 'Boarding:')
        response = self.search(first_names='')
        self.assertContains(response, 'Admitted students · 2028')
        self.assertNotContains(response, ROW[2])

    def test_family_can_search_and_view_six_relatives_with_default_limits(self):
        for number in range(6):
            name = f'Relative {number}'
            relative = AdmissionResult.objects.create(release=self.release, surname=ROW[1], first_names=name,
                primary_school='Family Primary', identity_key=identity_key(2027, ROW[1], name, 'Family Primary'))
            response = self.search(first_names=name)
            self.assertEqual(response.status_code, 200)
            self.assertContains(response, name)
            self.assertContains(self.view_result(relative), 'Boarding: Not offered')

    @override_settings(ADMISSION_SEARCH_RATE_PER_DAY=2)
    def test_repeated_search_limit_remains_configurable(self):
        self.assertEqual(self.search().status_code, 200)
        self.assertEqual(self.search().status_code, 200)
        self.assertEqual(self.search().status_code, 429)

    def test_partial_surname_confirms_admission_and_shows_school_links_without_boarding(self):
        AdmissionResult.objects.create(release=self.release, surname='Mokoena', first_names='Anna Jane',
            primary_school='Other School', identity_key=identity_key(2027, 'Mokoena', 'Anna Jane', 'Other School'), boarding=False)
        response = self.search(surname='koe', first_names='')
        self.assertContains(response, 'Mokoena, Thabo James')
        self.assertContains(response, 'Mokoena, Anna Jane')
        self.assertNotContains(response, 'Boarding:')
        self.assertContains(response, ROW[3])
        self.assertContains(response, 'Other School')
        self.assertContains(response, 'Admitted students ·')
        self.assertNotContains(response, 'name="primary_school"')
        self.assertNotContains(response, 'next steps')
        response = self.search(surname='moko', first_names='jam')
        self.assertContains(response, 'Mokoena, Thabo James')
        self.assertNotContains(response, 'Mokoena, Anna Jane')

    def test_school_only_empty_and_single_letter_surnames_cannot_search(self):
        for surname in ('', 'M', '%_'):
            response = self.search(surname=surname, first_names='', primary_school=ROW[3])
            self.assertNotContains(response, 'Admitted students ·')
            self.assertNotContains(response, ROW[2])

    def test_more_than_five_matches_requires_refinement_without_names_or_paging(self):
        for number in range(6):
            AdmissionResult.objects.create(release=self.release, surname='Mokoena', first_names=f'Student {number}',
                primary_school='Test Primary', identity_key=identity_key(2027, 'Mokoena', f'Student {number}', 'Test Primary'))
        response = self.search(surname='Moko', first_names='', page='2', offset='5')
        self.assertContains(response, 'Please narrow your search')
        self.assertNotContains(response, 'Student 0')
        self.assertNotContains(response, ROW[2])

    def test_school_spelling_is_display_only_and_click_targets_the_selected_student(self):
        other = AdmissionResult.objects.create(release=self.release, surname=ROW[1], first_names=ROW[2],
            primary_school='St Monicas', identity_key=identity_key(2027, ROW[1], ROW[2], 'St Monicas'), boarding=False)
        matches = self.search(surname='Mok', first_names='')
        candidate = next(item for item in matches.context['candidates'] if item['primary_school'] == other.primary_school)
        response = self.client.post(reverse('admission-results'), {
            'action': 'view', 'selection': candidate['selection']})
        self.assertContains(response, 'St Monicas')
        self.assertContains(response, 'Boarding: Not offered')
        self.assertNotContains(response, 'Boarding: Offered')
        response = self.view_result(self.result, primary_school="St Monica's", boarding='N')
        self.assertContains(response, ROW[3])
        self.assertContains(response, 'Boarding: Offered')

    def test_invalid_expired_and_stale_selection_tokens_do_not_reveal_result(self):
        token = selection_token(self.result)
        response = self.view_result(selection=token + 'invalid')
        self.assertNotContains(response, 'Boarding:')
        from django.core import signing
        with patch('school.admission_views.signing.loads', side_effect=signing.SignatureExpired()):
            self.assertNotContains(self.view_result(selection=token), 'Boarding:')
        import uuid
        self.release.revision = uuid.uuid4()
        self.release.save()
        self.assertNotContains(self.view_result(selection=token), 'Boarding:')

    def test_selection_is_unusable_after_withdrawal(self):
        token = selection_token(self.result)
        self.release.published = False
        self.release.save()
        self.assertNotContains(self.view_result(selection=token), 'Boarding:')

    def test_legacy_rows_need_reimport_and_disclaimer_removed(self):
        self.result.surname = ''
        self.result.first_names = ''
        self.result.save()
        response = self.search(surname='Moko', first_names='')
        self.assertNotContains(response, 'Admitted students ·')
        self.assertNotContains(response, 'Knowing someone')
        self.assertNotContains(response, 'name="own_details"')
        self.client.force_login(self.office)
        self.assertContains(self.client.get(reverse('staff-admission-results')), 'Re-upload this year')

    def test_old_name_only_imports_need_school_reimport_before_they_can_be_selected(self):
        self.result.primary_school = ''
        self.result.save()
        response = self.search(surname='Moko', first_names='')
        self.assertNotContains(response, 'Admitted students ·')
        self.assertNotContains(self.view_result(), 'Boarding:')
        self.client.force_login(self.office)
        self.assertContains(self.client.get(reverse('staff-admission-results')), 'Re-upload this year')

    def test_oversized_repeated_fields_and_file_uploads_rejected(self):
        path = reverse('admission-results')
        self.assertEqual(self.client.post(path, 'surname=A&surname=B', content_type='application/x-www-form-urlencoded').status_code, 400)
        self.assertEqual(self.client.post(path, {'file': upload()}).status_code, 400)
        self.assertEqual(self.client.post(path, 'A' * 70000, content_type='application/x-www-form-urlencoded').status_code, 413)

    def test_office_uploads_whole_year_hidden_and_publishes_separately(self):
        self.client.force_login(self.office)
        response = self.staff_upload([ROW, [2027, 'Example', 'Anna', 'Example Primary', 'N']])
        self.assertRedirects(response, reverse('staff-admission-results'))
        self.release.refresh_from_db()
        self.assertFalse(self.release.published)
        self.assertEqual(self.release.results.count(), 2)
        self.assertEqual(self.release.results.filter(boarding=True).count(), 1)
        self.assertFalse(AdmissionResult.objects.filter(pk=self.result.pk).exists())
        self.assertEqual(self.staff_action('publish').status_code, 302)
        self.assertContains(self.view_result(), 'Admitted.')
        self.assertEqual(self.staff_action('unpublish').status_code, 302)
        self.assertContains(self.search(), 'No matching students found')
        self.assertEqual(LogEntry.objects.filter(user=self.office).count(), 3)
        self.assertFalse(LogEntry.objects.filter(change_message__contains='Mokoena').exists())

    def test_invalid_upload_is_atomic_and_retains_published_results_and_other_year(self):
        self.client.force_login(self.office)
        other = AdmissionRelease.objects.create(year=2028, published=True)
        AdmissionResult.objects.create(release=other, identity_key=identity_key(2028, *ROW[1:4]))
        response = self.staff_upload([ROW, [2027, 'Broken', '', 'School', 'Y']])
        self.assertEqual(response.status_code, 400)
        self.release.refresh_from_db()
        self.assertTrue(self.release.published)
        self.assertTrue(AdmissionResult.objects.filter(pk=self.result.pk).exists())
        self.assertEqual(other.results.count(), 1)
        self.assertEqual(LogEntry.objects.filter(user=self.office).count(), 0)

    def test_replacement_only_affects_selected_year(self):
        self.client.force_login(self.office)
        self.assertEqual(self.staff_upload([[2028, *ROW[1:]]], 2028).status_code, 302)
        self.release.refresh_from_db()
        self.assertTrue(self.release.published)
        self.assertEqual(self.release.results.count(), 1)
        self.assertFalse(AdmissionRelease.objects.get(year=2028).published)

    def test_stale_publish_or_delete_cannot_affect_replaced_upload(self):
        self.client.force_login(self.office)
        stale_revision = str(self.release.revision)
        self.assertEqual(self.staff_upload().status_code, 302)
        for action in ('publish', 'delete'):
            response = self.staff_action(action, revision=stale_revision)
            self.assertEqual(response.status_code, 409)
            self.assertContains(response, 'changed after you opened', status_code=409)
        self.release.refresh_from_db()
        self.assertFalse(self.release.published)

    def test_delete_cascades_only_its_year_and_get_does_not_mutate(self):
        self.client.force_login(self.office)
        other = AdmissionRelease.objects.create(year=2028)
        response = self.client.get(reverse('staff-admission-results'), {'action': 'delete', 'year': 2027})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(AdmissionRelease.objects.filter(pk=self.release.pk).exists())
        self.assertEqual(self.staff_action('delete').status_code, 302)
        self.assertFalse(AdmissionResult.objects.filter(pk=self.result.pk).exists())
        self.assertTrue(AdmissionRelease.objects.filter(pk=other.pk).exists())

    def test_staff_permissions_and_navigation(self):
        path = reverse('staff-admission-results')
        self.assertEqual(self.client.get(path).status_code, 302)
        self.client.force_login(self.editor)
        self.assertEqual(self.client.get(path).status_code, 403)
        self.assertEqual(self.staff_upload().status_code, 403)
        self.client.force_login(self.office)
        response = self.client.get(path)
        self.assertContains(response, 'Admission results')
        self.assertContains(response, 'Withdraw results')
        self.assertNotContains(response, ROW[1])
        self.assertIn('no-store', response['Cache-Control'])
        readonly = User.objects.create_user('results-viewer', is_staff=True)
        readonly.user_permissions.add(Permission.objects.get(codename='view_admissionrelease'))
        self.client.force_login(readonly)
        self.assertEqual(self.client.get(path).status_code, 200)
        self.assertEqual(self.staff_upload().status_code, 403)
        self.assertEqual(self.staff_action('publish').status_code, 403)
        self.assertEqual(self.staff_action('delete').status_code, 403)

    def test_template_download_is_private_and_contains_only_headers(self):
        path = reverse('staff-admission-results-template')
        self.assertEqual(self.client.get(path).status_code, 302)
        self.client.force_login(self.editor)
        self.assertEqual(self.client.get(path).status_code, 403)
        self.client.force_login(self.office)
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content.decode().strip(), ','.join(HEADERS))
        self.assertIn('no-store', response['Cache-Control'])

    def test_database_failure_rolls_back_bulk_replacement(self):
        self.client.force_login(self.office)
        with patch('school.admission_views.AdmissionResult.objects.bulk_create', side_effect=RuntimeError('Test failure')):
            with self.assertRaises(RuntimeError):
                self.staff_upload()
        self.release.refresh_from_db()
        self.assertTrue(self.release.published)
        self.assertTrue(AdmissionResult.objects.filter(pk=self.result.pk).exists())
        self.assertFalse(LogEntry.objects.filter(user=self.office).exists())

    def test_unexpected_error_reports_redact_submitted_details_and_view_locals(self):
        from .admission_views import lookup
        request = RequestFactory().post(reverse('admission-results'), {
            'year': ROW[0], 'surname': ROW[1], 'first_names': ROW[2],
            'primary_school': ROW[3], 'own_details': 'on'})
        with patch('school.admission_views.consume_limit', side_effect=RuntimeError('Deliberate test error')):
            try:
                lookup(request)
            except RuntimeError:
                exc_type, exc_value, tb = sys.exc_info()
                # Start at the view stack, as a server error handler would.
                report = ExceptionReporter(request, exc_type, exc_value, tb.tb_next).get_traceback_text()
            else:
                self.fail('Expected the deliberately injected exception')
        for value in ROW[1:4]:
            self.assertNotIn(value, report)

    def test_invalid_action_and_malicious_values(self):
        self.client.force_login(self.office)
        self.assertEqual(self.client.post(reverse('staff-admission-results'), {'action': 'arbitrary'}).status_code, 400)
        response = self.search(surname='<script>alert(1)</script>')
        self.assertNotContains(response, '<script>alert(1)</script>')

    def test_lookup_link_available_when_applications_closed_and_no_indexing(self):
        self.assertContains(self.client.get('/apply/'), reverse('admission-results'))
        self.assertNotContains(self.client.get('/sitemap.xml'), reverse('admission-results'))
        self.assertContains(self.client.get('/robots.txt'), 'Disallow: /admissions/results/')
        response = self.client.get(reverse('admission-results'))
        self.assertContains(response, 'noindex, nofollow')
