from datetime import date
from io import BytesIO

from django.contrib.admin.models import LogEntry
from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.db import models
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from .models import Application, SiteSettings
from .staff_forms import AdmissionsSettingsForm


@override_settings(SECURE_SSL_REDIRECT=False, STORAGES={
    'default': {'BACKEND': 'django.core.files.storage.InMemoryStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
})
class AdmissionsWorkspaceTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('setup_staff_groups', verbosity=0)
        cls.site = SiteSettings.objects.create(admissions_open=False)
        cls.office = User.objects.create_user('admissions-office', is_staff=True)
        cls.office.groups.add(Group.objects.get(name='Admissions officers'))
        cls.editor = User.objects.create_user('admissions-editor', is_staff=True)
        cls.editor.groups.add(Group.objects.get(name='Content editors'))
        cls.ordinary = User.objects.create_user('ordinary-admissions')
        for number in range(23):
            Application.objects.create(learner_name=f'Learner {number}', date_of_birth=date(2013, 1, 1),
                entry_level='Grade 8', application_year=2027, guardian_name='Guardian Example',
                phone='0123456789', email='guardian@example.com', consent=True,
                status='reviewing', staff_notes='Never export this private note')
        cls.other = Application.objects.create(learner_name='Separate Applicant', date_of_birth=date(2012, 2, 1),
            entry_level='Grade 8', guardian_name='Other Guardian', phone='55500000', consent=True)

    def test_closed_admissions_hide_all_intake_details_and_form_even_after_post(self):
        for response in [self.client.get('/apply/'), self.client.post('/apply/', {'surname': 'Example'})]:
            self.assertContains(response, 'Admissions are currently closed')
            for text in ['Application fee', 'Payment method', 'Interview day', 'What to bring',
                         'admissions-facts', 'name="surname"', 'Your first step towards joining']:
                self.assertNotContains(response, text)
        self.assertEqual(Application.objects.count(), 24)
        response = self.client.get('/')
        self.assertContains(response, '>Admissions</a>')
        self.assertContains(response, 'href="/admissions/results/"')

    def test_office_can_update_only_admissions_settings(self):
        self.client.force_login(self.office)
        self.assertEqual(self.client.get('/staff/settings/').status_code, 403)
        response = self.client.get('/staff/admissions/settings/')
        self.assertContains(response, 'Admissions settings')
        self.assertContains(response, 'name="admissions_fee"')
        self.assertNotContains(response, 'name="hero_title"')
        payload = {name: getattr(self.site, name) for name in AdmissionsSettingsForm.Meta.fields
                   if not isinstance(self.site._meta.get_field(name), models.ImageField)}
        payload = {name: '' if value is None else value for name, value in payload.items()}
        payload.update(admissions_open='on', admissions_year=2029, admissions_fee='70.00', name='Malicious school rename')
        response = self.client.post('/staff/admissions/settings/', payload)
        self.assertRedirects(response, '/staff/admissions/settings/')
        self.site.refresh_from_db()
        self.assertTrue(self.site.admissions_open)
        self.assertEqual(self.site.admissions_year, 2029)
        self.assertEqual(self.site.name, 'Holy Family High School')
        self.assertTrue(LogEntry.objects.filter(change_message__contains='Admissions settings updated').exists())

    def test_general_settings_no_longer_render_admissions_fields(self):
        self.client.force_login(self.editor)
        response = self.client.get('/staff/settings/')
        self.assertNotContains(response, 'name="admissions_open"')
        self.assertNotContains(response, 'name="admissions_fee"')
        self.assertContains(response, '/staff/admissions/settings/')

    def workbook(self, path):
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'no-store, private')
        self.assertIn('.xlsx', response['Content-Disposition'])
        data = b''.join(response.streaming_content)
        response.close()
        return load_workbook(BytesIO(data))

    def test_excel_exports_every_filtered_record_and_excludes_internal_notes(self):
        self.client.force_login(self.office)
        workbook = self.workbook('/staff/applications/export/?state=reviewing&q=Learner&page=2')
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 24)  # Header plus all 23 matches, independent of pagination.
        self.assertEqual(sheet.freeze_panes, 'A2')
        rows = list(sheet.values)
        self.assertNotIn('Never export this private note', str(rows))
        self.assertNotIn('Separate Applicant', str(rows))
        self.assertTrue(any('23 application records' in entry.change_message for entry in LogEntry.objects.all()))
        self.assertEqual(sheet['M2'].value, '0123456789')
        self.assertEqual(sheet['E2'].value.date(), date(2013, 1, 1))
        workbook.close()

    def test_excel_neutralises_formula_values_and_empty_export_is_valid(self):
        Application.objects.filter(pk=self.other.pk).update(current_school='=HYPERLINK("https://example.com")', message='+1+1')
        self.client.force_login(self.office)
        workbook = self.workbook('/staff/applications/export/?q=Separate')
        sheet = workbook.active
        self.assertEqual(sheet['J2'].data_type, 's')
        self.assertEqual(sheet['J2'].value, '=HYPERLINK("https://example.com")')
        self.assertEqual(sheet['O2'].data_type, 's')
        workbook.close()
        workbook = self.workbook('/staff/applications/export/?q=nonexistent-record')
        self.assertEqual(workbook.active.max_row, 1)
        workbook.close()

    def test_private_export_and_settings_require_authorised_staff(self):
        for user in [None, self.ordinary, self.editor]:
            self.client.logout()
            if user:
                self.client.force_login(user)
            response = self.client.get('/staff/applications/export/')
            self.assertEqual(response.status_code, 302 if user is None else 403)
        self.client.force_login(self.office)
        self.assertEqual(self.client.post('/staff/applications/export/').status_code, 405)
