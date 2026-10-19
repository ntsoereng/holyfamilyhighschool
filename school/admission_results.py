"""Private bulk admission imports, partial name search and selected results."""
import csv
from io import StringIO
import json
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException
from django import forms
from django.core.exceptions import ValidationError
from django.core import signing
from django.utils.crypto import salted_hmac
from django.views.decorators.debug import sensitive_variables

from .submissions import normalize_name, normalize_text

HEADERS = ('Year', 'Surname', 'First Name(s)', 'Primary School', 'Boarding')
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_UNPACKED_BYTES = 20 * 1024 * 1024
MAX_ROWS = 5000
MAX_MATCHES = 5
SELECTION_SALT = 'admission-result-selection-v1'
SELECTION_MAX_AGE = 600


def selection_token(result):
    return signing.dumps({'result': result.pk, 'year': result.release.year,
                          'revision': str(result.release.revision)}, salt=SELECTION_SALT)


@sensitive_variables()
def identity_key(year, surname, first_names, primary_school):
    parts = [int(year), *(normalize_name(value).casefold() for value in (surname, first_names, primary_school))]
    return salted_hmac('admission-result-v1', json.dumps(parts, ensure_ascii=False), algorithm='sha256').hexdigest()


class LookupForm(forms.Form):
    surname = forms.CharField(max_length=80, help_text='Enter at least two letters. Part of a surname is enough.')
    first_names = forms.CharField(label='First name(s)', max_length=120, required=False,
                                 help_text='Add a first name or part of it to narrow the matches.')
    website = forms.CharField(required=False, widget=forms.HiddenInput)

    def clean(self):
        data = super().clean()
        for name in ('surname', 'first_names'):
            if name in data:
                data[name] = normalize_name(data[name])
        if 'surname' in data and sum(char.isalpha() for char in data['surname']) < 2:
            self.add_error('surname', 'Enter at least two letters of the surname.')
        if data.get('website'):
            raise ValidationError('We could not verify these details. Please contact the school office.')
        return data


class SelectionForm(forms.Form):
    selection = forms.CharField(max_length=500, widget=forms.HiddenInput)


def _xlsx_rows(upload):
    # Bound the compressed input and archive expansion before openpyxl loads it.
    try:
        with ZipFile(upload) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(entry.file_size for entry in entries) > MAX_UNPACKED_BYTES:
                raise ValidationError('This workbook is too large when unpacked. Use a simple values-only workbook.')
            for entry in entries:
                if entry.flag_bits & 1 or entry.filename.lower().endswith('vbaproject.bin'):
                    raise ValidationError('Encrypted workbooks and macros are not supported.')
                if entry.filename.lower().endswith(('.xml', '.rels')):
                    # Preflight every XML part, even if openpyxl would ignore it.
                    ElementTree.fromstring(archive.read(entry), forbid_dtd=True)
        upload.seek(0)
        from openpyxl import load_workbook
        workbook = load_workbook(upload, read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValidationError('Use a workbook with exactly one worksheet.')
            sheet = workbook.worksheets[0]
            # Do not trust worksheet dimensions supplied by the file.
            sheet.reset_dimensions()
            for number, row in enumerate(sheet.iter_rows(), 1):
                if number > MAX_ROWS + 1 or len(row) > 5:
                    raise ValidationError(f'Use only five columns and no more than {MAX_ROWS} students.')
                if any(cell.data_type in {'f', 'e'} for cell in row):
                    raise ValidationError(f'Row {number}: use values only, without formulas or Excel errors.')
                yield [cell.value for cell in row]
        finally:
            workbook.close()
    except ValidationError:
        raise
    except (BadZipFile, DefusedXmlException, ElementTree.ParseError, KeyError, ValueError, TypeError, OSError, EOFError) as exc:
        raise ValidationError('This is not a readable, safe .xlsx workbook. Save a new values-only workbook and try again.') from exc


@sensitive_variables()
def read_results(upload, year):
    if upload.size > MAX_UPLOAD_BYTES:
        raise ValidationError('Choose a spreadsheet smaller than 2 MB.')
    extension = Path(upload.name).suffix.lower()
    if extension == '.xlsx':
        rows = _xlsx_rows(upload)
    elif extension == '.csv':
        try:
            rows = csv.reader(StringIO(upload.read(MAX_UPLOAD_BYTES + 1).decode('utf-8-sig')), strict=True)
        except UnicodeDecodeError as exc:
            raise ValidationError('Save the CSV using UTF-8 encoding.') from exc
    else:
        raise ValidationError('Choose an Excel .xlsx file or a UTF-8 .csv file. Older .xls files are not supported.')
    results = []
    seen = set()
    try:
        header = next(rows, None)
        if not header or [normalize_text(str(value or '')).casefold() for value in header] != [value.casefold() for value in HEADERS]:
            raise ValidationError('The first row must be: Year, Surname, First Name(s), Primary School, Boarding.')
        for number, row in enumerate(rows, 2):
            if number > MAX_ROWS + 1:
                raise ValidationError(f'Upload no more than {MAX_ROWS} student rows.')
            if not any(value is not None and str(value).strip() for value in row):
                continue
            if len(row) > 5:
                raise ValidationError(f'Row {number}: use exactly five columns.')
            row = list(row) + [None] * (5 - len(row))
            row_year, surname, first_names, primary_school, boarding = row
            if isinstance(row_year, bool) or str(row_year).strip() != str(year):
                raise ValidationError(f'Row {number}: every Year must match the selected academic year ({year}).')
            names = []
            for value, maximum in zip((surname, first_names, primary_school), (80, 120, 180)):
                if not isinstance(value, str):
                    raise ValidationError(f'Row {number}: surname, first names and primary school must be text.')
                value = normalize_name(value)
                if not value or len(value) > maximum or value.startswith(('=', '+', '-', '@')):
                    raise ValidationError(f'Row {number}: check the required names and primary school (maximum lengths 80, 120 and 180).')
                names.append(value)
            boarding = str('' if boarding is None else boarding).strip().upper()
            if boarding not in {'Y', 'N', ''}:
                raise ValidationError(f'Row {number}: Boarding must be Y, N or blank.')
            key = identity_key(year, *names)
            if key in seen:
                raise ValidationError(f'Row {number}: duplicate student details. Resolve duplicates before uploading.')
            seen.add(key)
            results.append({'identity_key': key, 'surname': names[0], 'first_names': names[1],
                            'primary_school': normalize_text(primary_school), 'boarding': boarding == 'Y'})
    except (csv.Error, UnicodeDecodeError) as exc:
        raise ValidationError('The CSV could not be read. Use a UTF-8 CSV with five columns.') from exc
    if not results:
        raise ValidationError('The spreadsheet contains no students. Existing results have not been changed.')
    return results


class ImportForm(forms.Form):
    year = forms.IntegerField(label='Academic year', min_value=2000, max_value=2100)
    spreadsheet = forms.FileField(label='Admission spreadsheet', widget=forms.FileInput(attrs={'accept': '.xlsx,.csv'}))
    replace = forms.BooleanField(label='Replace this year’s results with this upload and hide them until I publish again.')

    def clean(self):
        data = super().clean()
        if 'spreadsheet' in data and 'year' in data:
            data['rows'] = read_results(data['spreadsheet'], data['year'])
        return data


class ReleaseActionForm(forms.Form):
    year = forms.IntegerField(min_value=2000, max_value=2100)
    revision = forms.UUIDField()
    action = forms.ChoiceField(choices=[('publish', 'Publish'), ('unpublish', 'Unpublish'), ('delete', 'Delete')])
