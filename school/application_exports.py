"""Typed Excel exports; submitted text is never written as a spreadsheet formula."""
from tempfile import SpooledTemporaryFile

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = (
    ('Reference', 38), ('Given name', 22), ('Surname', 22), ('Learner name', 30),
    ('Date of birth', 16), ('Entry level', 16), ('Intake year', 14), ('Address', 40),
    ('District', 22), ('Current school', 30), ('Guardian name', 30), ('Email', 34),
    ('Phone', 22), ('Boarding required', 20), ('Message', 45), ('Consent provided', 20),
    ('Review status', 18), ('Submitted (Lesotho time)', 25),
)


def export_applications(records):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('Applications')
    sheet.freeze_panes = 'A2'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    for number, (_, width) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(number)].width = width

    def cell(value, header=False):
        result = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            result.value = ILLEGAL_CHARACTERS_RE.sub('', value)
            result.data_type = 's'  # Includes strings beginning with =, +, - or @.
        if header:
            result.font = Font(bold=True, color='FFFFFF')
            result.fill = PatternFill('solid', fgColor='174B3B')
        elif hasattr(value, 'hour'):
            result.number_format = 'yyyy-mm-dd hh:mm'
        elif hasattr(value, 'year'):
            result.number_format = 'yyyy-mm-dd'
        return result

    sheet.append([cell(title, header=True) for title, _ in COLUMNS])
    count = 0
    for item in records.iterator(chunk_size=1000):
        values = (
            item.reference, item.given_name, item.surname, item.learner_name,
            item.date_of_birth, item.entry_level, item.application_year, item.physical_address,
            item.district_display, item.current_school, item.guardian_name, item.email,
            item.phone, 'Not provided' if item.boarding_required is None else 'Yes' if item.boarding_required else 'No',
            item.message, 'Yes' if item.consent else 'No', item.get_status_display(),
            timezone.localtime(item.submitted_at).replace(tzinfo=None),
        )
        sheet.append([cell(value) for value in values])
        count += 1
    sheet.auto_filter.ref = f'A1:R{count + 1}'
    output = SpooledTemporaryFile(max_size=5 * 1024 * 1024, mode='w+b')
    try:
        workbook.save(output)
        output.seek(0)
    except Exception:
        output.close()
        raise
    return output, count
